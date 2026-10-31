import io
import os

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from pptx import Presentation


class DocumentProcessor:
    @staticmethod
    def create_document_bytes(file_format):
        output = io.BytesIO()
        if file_format == "txt":
            return b""
        if file_format == "docx":
            Document().save(output)
        elif file_format == "xlsx":
            from openpyxl import Workbook
            Workbook().save(output)
        elif file_format == "pptx":
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(914400, 914400, 7315200, 1828800)
            box.text_frame.text = "Введите текст слайда"
            presentation.save(output)
        else:
            raise ValueError("Неизвестный формат документа")
        return output.getvalue()

    @staticmethod
    def detect_format(file_name, mime_type=""):
        extension = os.path.splitext(str(file_name).lower())[1]
        by_extension = {
            ".txt": "txt",
            ".docx": "docx",
            ".xlsx": "xlsx",
            ".pdf": "pdf",
            ".pptx": "pptx",
        }
        if extension in by_extension:
            return by_extension[extension]
        by_mime = {
            "text/plain": "txt",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
            "application/pdf": "pdf",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
        }
        if mime_type in by_mime:
            return by_mime[mime_type]
        raise ValueError(
            "Формат не поддерживается. Используйте TXT, DOCX, XLSX, PPTX или PDF."
        )

    @staticmethod
    def read_txt_bytes(file_bytes):
        return file_bytes.decode("utf-8-sig", errors="replace")

    @staticmethod
    def modify_txt_bytes(new_text):
        return new_text.encode("utf-8")

    @staticmethod
    def _iter_word_paragraphs(parent, seen_cells=None):
        if seen_cells is None:
            seen_cells = set()
        if hasattr(parent, "element") and hasattr(parent.element, "body"):
            container = parent.element.body
        elif hasattr(parent, "_tc"):
            container = parent._tc
        else:
            container = parent
        for child in container.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, parent)
            elif child.tag == qn("w:tbl"):
                table = Table(child, parent)
                for row in table.rows:
                    for cell in row.cells:
                        cell_key = id(cell._tc)
                        if cell_key in seen_cells:
                            continue
                        seen_cells.add(cell_key)
                        yield from DocumentProcessor._iter_word_paragraphs(
                            cell, seen_cells
                        )

    @staticmethod
    def read_word_bytes(file_bytes):
        document = Document(io.BytesIO(file_bytes))
        paragraphs = list(DocumentProcessor._iter_word_paragraphs(document))
        return "\n".join(paragraph.text for paragraph in paragraphs)

    @staticmethod
    def _replace_paragraph_text_preserving_xml(paragraph, value):
        text_nodes = paragraph._p.xpath(".//w:t")
        if text_nodes:
            text_nodes[0].text = value
            for node in text_nodes[1:]:
                node.text = ""
        elif value:
            paragraph.add_run(value)

    @staticmethod
    def modify_word_bytes(old_bytes, new_text):
        # Keep the original document package. Replace only text nodes, so
        # paragraph properties, tables, pictures, headers, and other XML parts
        # remain in the DOCX instead of rebuilding a blank document.
        document = Document(io.BytesIO(old_bytes))
        paragraphs = list(DocumentProcessor._iter_word_paragraphs(document))
        lines = new_text.split("\n")

        for index, paragraph in enumerate(paragraphs):
            value = lines[index] if index < len(lines) else ""
            DocumentProcessor._replace_paragraph_text_preserving_xml(
                paragraph, value
            )

        for value in lines[len(paragraphs):]:
            document.add_paragraph(value)

        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    @staticmethod
    def read_excel_bytes(file_bytes):
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=False)
        try:
            sheet = workbook.active
            rows = []
            for row in sheet.iter_rows():
                values = [cell.value for cell in row]
                if any(value is not None for value in values):
                    rows.append(" | ".join(
                        "" if value is None else str(value) for value in values
                    ))
            return "\n".join(rows)
        finally:
            workbook.close()

    @staticmethod
    def _parse_excel_value(value):
        value = value.strip()
        if value == "":
            return None
        try:
            if value.lstrip("+-").isdigit():
                return int(value)
            return float(value.replace(",", "."))
        except ValueError:
            return value

    @staticmethod
    def modify_excel_bytes(old_bytes, new_text):
        # Update only the active sheet's values; keep workbook sheets,
        # formatting, widths, merged cells, and other workbook structures.
        workbook = load_workbook(io.BytesIO(old_bytes), data_only=False)
        try:
            sheet = workbook.active
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.__class__.__name__ != "MergedCell":
                        cell.value = None

            for row_index, line in enumerate(new_text.split("\n"), start=1):
                for column_index, value in enumerate(line.split(" | "), start=1):
                    target = sheet.cell(row=row_index, column=column_index)
                    if target.__class__.__name__ != "MergedCell":
                        target.value = DocumentProcessor._parse_excel_value(value)

            output = io.BytesIO()
            workbook.save(output)
            return output.getvalue()
        finally:
            workbook.close()

    @staticmethod
    def read_pdf_bytes(file_bytes):
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    @staticmethod
    def modify_pdf_bytes(old_bytes, new_text):
        # pypdf cannot rewrite arbitrary page text in-place. Preserve every
        # page and store the edited text in PDF Description metadata.
        reader = PdfReader(io.BytesIO(old_bytes))
        writer = PdfWriter()
        writer.append(reader)
        metadata = dict(reader.metadata or {})
        metadata["/Description"] = new_text
        writer.add_metadata(metadata)
        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()


    @staticmethod
    def read_pptx_bytes(file_bytes):
        presentation = Presentation(io.BytesIO(file_bytes))
        lines = []
        for number, slide in enumerate(presentation.slides, start=1):
            lines.append("[Слайд " + str(number) + "]")
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    lines.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)

    @staticmethod
    def modify_pptx_bytes(old_bytes, new_text):
        presentation = Presentation(io.BytesIO(old_bytes))
        paragraphs = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    paragraphs.extend(shape.text_frame.paragraphs)
        lines = [line for line in new_text.split("\n") if not line.strip().startswith("[Слайд ")]
        for index, paragraph in enumerate(paragraphs):
            value = lines[index] if index < len(lines) else ""
            if paragraph.runs:
                paragraph.runs[0].text = value
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = value
        output = io.BytesIO()
        presentation.save(output)
        return output.getvalue()
